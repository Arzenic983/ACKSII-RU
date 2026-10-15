import collections
import json
import pathlib
import sys
from zipfile import ZipFile
from lxml import etree
import openpyxl

sys.stdout.reconfigure(encoding='utf8')
ROOT=pathlib.Path(__file__).resolve().parent.parent
OUT=ROOT/'.work'/'translation-database'
OUT.mkdir(exist_ok=True)
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def paragraph(p):
    return ''.join(p.itertext()) if False else ''.join(p.xpath('.//w:t/text() | .//w:tab/@w:val',namespaces=NS))

def docx(path):
    with ZipFile(path) as z:xml=etree.fromstring(z.read('word/document.xml'))
    items=[]
    for i,node in enumerate(xml.find('w:body',NS)):
        kind=etree.QName(node).localname
        if kind=='p':
            text=paragraph(node)
            if not text.strip():continue
            style=node.find('w:pPr/w:pStyle',NS)
            items.append({'locator':f'body.{i+1}','kind':'paragraph','style':style.get('{'+NS['w']+'}val') if style is not None else None,'text':text})
        elif kind=='tbl':
            rows=[]
            for r in node.findall('w:tr',NS):
                row=[]
                for c in r.findall('w:tc',NS):
                    row.append('\n'.join(paragraph(p) for p in c.findall('w:p',NS)))
                rows.append(row)
            items.append({'locator':f'body.{i+1}','kind':'table','rows':rows})
    return {'path':str(path.relative_to(ROOT)),'items':items}

paths=[ROOT/'ACKSII'/'ReadMe.docx',ROOT/'ACKSII'/'ACKSII Revised Rulebook'/'CHAPTER 1_ CHARACTERS - Персонажи.docx']
for path in paths:
    data=docx(path)
    name='readme' if path.name=='ReadMe.docx' else 'existing_chapter_01'
    (OUT/(name+'.json')).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
    lines=[]
    for item in data['items']:
        lines.append(f"[{item['locator']}] {item['kind']} {item.get('style') or ''}")
        if item['kind']=='paragraph':lines.append(item['text'])
        else:lines.extend(' | '.join(row) for row in item['rows'])
        lines.append('')
    (OUT/(name+'.txt')).write_text('\n'.join(lines),encoding='utf8')
    print(name,'items',len(data['items']),'chars',len('\n'.join(lines)))
book=openpyxl.load_workbook(ROOT/'ACKSII'/'Прогресс перевода.xlsx',data_only=False,read_only=True)
sheets={s.title:[list(r) for r in s.iter_rows(values_only=True)] for s in book}
(OUT/'progress_snapshot.json').write_text(json.dumps(sheets,ensure_ascii=False,indent=2,default=str),encoding='utf8')
print('sheets',[(s.title,s.max_row,s.max_column) for s in book])
