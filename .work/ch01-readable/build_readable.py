"""Assemble a separate reading copy, without writing to translation inputs."""
import collections
import hashlib
import html
import json
import pathlib
import re
import sqlite3
import sys
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

sys.stdout.reconfigure(encoding='utf8')
WORK=pathlib.Path(__file__).resolve().parent
ROOT=WORK.parents[1]
DB=ROOT/'ACKSII'/'База перевода'/'acksii_translation.sqlite'
OUT=ROOT/'ACKSII'/'Сборки для чтения'
OUT.mkdir(exist_ok=True)
FILE=OUT/'ACKSII Введение и глава 1 с пометками.docx'
BASELINE=WORK/'input_hashes.json'
if not BASELINE.exists():
    protected=[]
    for pattern in ('*.docx','*.xlsx','*.sqlite','*.md','*.json','*.jsonl','*.pdf'):
        protected.extend(p for p in (ROOT/'ACKSII').rglob(pattern) if OUT not in p.parents)
    hashes={p.relative_to(ROOT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(set(protected))}
    BASELINE.write_text(json.dumps(hashes,ensure_ascii=False,indent=2),encoding='utf8')

db=sqlite3.connect(DB.as_uri()+'?mode=ro',uri=True)
db.row_factory=sqlite3.Row
units=[dict(r) for r in db.execute('SELECT * FROM bilingual WHERE pdf_page<=23 ORDER BY pdf_page,block_order,row_index,column_index,id')]
unit_map={x['id']:x for x in units}
tables={r['id']:dict(r) for r in db.execute('SELECT * FROM source_tables WHERE pdf_page<=23')}
assert len(units)==1689 and all(u['russian'] for u in units)

# Every report item has at least one explicit anchor. A question is labelled as
# a question and never masquerades as a correction that has already been made.
notes=[]
def note(uid,text,category='Исправление',key=None,part=None):
    assert uid in unit_map,uid
    notes.append({'unit_id':uid,'text':text,'category':category,'report_key':key,'part':part})

note('rr.p0002.b006','На обложке оригинала указано 18 классов, а в оглавлении и главе 2 перечислен 21. Число 18 сохранено; это расхождение оригинала, а не внесённая правка.','Вопрос к оригиналу','ambiguity_classes',0)
note('rr.p0005.t01.r063.c01','В словаре стоит «Нобирийский чудотворцец?». Написание сохранено. Предложение из отчёта: «Нобирийский чудотворец». Исправление словаря пока не применено.','Вопрос к словарю','ambiguity_nobiran')
note('rr.p0005.t01.r057.c01','В словаре указаны варианты «Ведьма / Ворожей». Здесь используется первая форма; выбор альтернативы не считается утверждённым.','Вопрос к словарю','ambiguity_variants')
note('rr.p0005.t01.r153.c01','Сохранена словарная форма «Фокусничество». В ReadMe она сопровождается вопросительным знаком и предложениями «Волшебные фокусы», «Фокусы», «Фокусник». Новая замена не внесена.','Вопрос к словарю','ambiguity_variants')
note('rr.p0005.t01.r158.c01','Сохранено словарное «Реликварианизм». В ReadMe есть редакционная помета с вариантом «Почитание реликвий»; он пока не применён.','Вопрос к словарю','ambiguity_variants')
note('rr.p0005.t01.r134.c01','Утверждённая пользователем форма Loremastery — «Мифография», вместо «Знание преданий».','Терминология','terminology')
note('rr.p0005.t01.r147.c01','Утверждённая пользователем форма Naturalism — «Природоведение», вместо «Знание природы».','Терминология','terminology')
note('rr.p0006.t01.r066.c01','Spells как название главы — «Заклинания». Словарное «от заклинаний» относится к спасброскам и к этому заголовку не применяется.','Терминология','terminology')
note('rr.p0007.t01.r158.c01','В этом разделе оглавления перечислены народы. Поэтому Skysos — «Скизийцы» по полной словарной записи; «Скизия» означает страну.','Исправление','skysos')
note('rr.p0009.b023','Во всей главе роль Judge обозначена словарным «Арбитр». Разнородные «Судья» и «мастер» в черновике приведены к этой форме. PC — «ПИ», NPC — «НИП».','Терминология','terminology')
note('rr.p0010.b007','Термины различаются по словарю: Domain — «Домен», Realm — «Владение». В переводе они не подменяют друг друга.','Терминология','terminology')
note('rr.p0013.t02.r001.c03','Extraordinary damage передаётся как «Особый урон» — форма, согласованная с пользователем. Категория урона сохранена отдельно от его типа.','Терминология','terminology')
note('rr.p0016.b001','Названия и сокращения атрибутов приведены к словарю: Сила — СИЛ, Интеллект — ИНТ, Воля — ВОЛ, Ловкость — ЛОВ, Телосложение — ТЕЛ, Харизма — ХАР.','Терминология','terminology')
note('rr.p0016.b013','Этот абзац отсутствовал в исходном переводе. Восстановлены требования: не менее 9 в каждом ключевом атрибуте, а также выполнение дополнительных минимальных требований класса.',key='r01')
note('rr.p0016.b008','В черновике действие Ловкости было ограничено метательным оружием. Missile weapons включает всё дальнобойное оружие; соответствующий бонус или штраф действует на все такие броски атаки.',key='r02')
note('rr.p0016.b010','Лимит Харизмы относится к приспешникам, а не к наёмникам. Сохранена формула: 4 плюс модификатор Харизмы, то есть от 1 до 7 приспешников.',key='r03')
note('rr.p0017.b023','В черновике был неверно указан источник теряемых бонусов. При таком снаряжении не применяются бонусы атрибутов к броскам атаки и КБ. Ограничения классовых способностей и получения опыта также сохранены.',key='r04')
note('rr.p0018.b001','Использованы словарные «очки здоровья (ОЗ)» и «кости здоровья (КЗ)» вместо разнородных обозначений черновика.','Терминология','terminology')
note('rr.p0018.b004','Исправлено ошибочное «переписывание» ОЗ: при повышении уровня перебрасывают все КЗ. При результате не выше прежнего добавляют 1 ОЗ. Начиная с 10-го уровня новых КЗ нет, но полученные ранее КЗ всё равно перебрасывают.',key='r05')
note('rr.p0018.b008','Venturer — класс «Торговец», а не общее обозначение искателя приключений. Поэтому в примере максимальных ОЗ указаны торговцы, 4 ОЗ до модификатора Телосложения.',key='r06')
note('rr.p0019.b017','Скорость исследования измеряется в футах за ход. Ошибочное «за раунд» в формуле таблицы черновика исправлено; скорость ненагруженного человека остаётся 120 футов за ход.',key='r07')
note('rr.p0019.t01.r002.c02','В таблице черновика было 30 футов за раунд. Исправлено на 40 футов за раунд, как в английском PDF.',key='r07')
note('rr.p0019.t01.r003.c01','В черновике здесь стояли футы за раунд. Исправлено: футы за ход. Соотношения четырёх скоростей соответствуют оригиналу.',key='r07')
note('rr.p0019.b023','Сохранены исходные фунты и стоуны: исторически 8–14 фунтов за стоун, в ACKS II условно около 10 фунтов. Добавленные в черновике метрические пересчёты убраны.',key='r08')
note('rr.p0020.b005','Тяжёлые предметы включают предметы весом 8 фунтов и больше. Ошибочный добавленный пересчёт в 1,8 кг удалён; в сборке используются единицы оригинала.',key='r08')
note('rr.p0020.b004','Восстановлены исходные примеры комплектов: 6 железных шипов и 20 стрел. Черновик подменял состав и количество предметов.',key='r09')
note('rr.p0020.b009','Iron rations переведено как «сухой паёк» вместо буквального «железный паёк». Расчёт нагрузки сохранён: 12 обычных предметов, 2 тяжёлых предмета и доспехи дают 10 стоунов.',key='r09')
note('rr.p0020.b015','Восстановлено утраченное правило грамотности: персонаж с Интеллектом не ниже среднего умеет читать и писать на всех языках, на которых говорит.',key='r10')
note('rr.p0021.b004','Черновик утверждал, что эльфы и нобирийцы не достигают взрослого возраста. Исправлено: они достигают его, но не стареют дальше этой категории.',key='r11')
note('rr.p0021.t02.r001.c02','Youth — «Юность». Ошибочное «Детство» в черновике заменено словарной возрастной категорией.',key='r11')
note('rr.p0021.t02.r001.c03','Adult — «Взрослый». Ошибочное «Подросток» в черновике заменено словарной возрастной категорией.',key='r11')
note('rr.p0022.b004','Категория размера отражает массу и объём и не равна конкретному росту. Все доступные расы имеют размер «с человека»; обычный вес — 15 стоунов.',key='r12')
note('rr.p0023.b007','Восстановлено пропущенное правило: продолжительное занятие обычно можно прервать без штрафа или нарушения хода работы.',key='r13')
note('rr.p0023.b009','Форсированный марш увеличивает время в пути и расстояние вплоть до 50%, а не обязательно ровно на 50%. Сохранены запрет вспомогательных занятий и немедленная потребность в отдыхе.',key='r13')
note('rr.p0019.b016','Словарь переводит Traversing как «прорывание»; форма сохранена. По описанию оригинала речь о переходах через опасные участки, в том числе по верёвке над пропастью. Предложение из отчёта — «Переправа» — пока не применено.','Вопрос к словарю','ambiguity_traversing')
note('rr.p0020.b011','Ссылка на печатную стр. 487 оставлена по оригиналу: там начинается текст Приложения A. Его заставка находится на стр. 486. Физические страницы PDF — 489 и 488 соответственно.','Вопрос к оригиналу','ambiguity_pages')

# The table contains the most compact complete set of class names in chapter 1.
draft=json.loads(db.execute("SELECT metadata_json FROM draft_fragments WHERE source_id='ch01_draft' AND locator='body.149'").fetchone()[0])['rows']
old_by_class={r['id']:r['raw_text'] for r in db.execute('''SELECT u.id,f.raw_text FROM units u JOIN unit_drafts ud ON ud.unit_id=u.id
 JOIN draft_fragments f ON f.id=ud.draft_id WHERE u.table_id='rr.p0021.t01' AND u.column_index%2=0 AND u.row_index>0''')}
changes=[]
for uid,old in old_by_class.items():
    new=unit_map[uid]['russian']
    if old.casefold()!=new.casefold():changes.append(f'«{old}» → «{new}»')
note('rr.p0021.t01.r001.c01','Названия классов приведены к словарю проекта: '+ '; '.join(changes)+'.','Терминология','terminology')
required={f'r{i:02d}' for i in range(1,14)}|{'terminology','skysos','ambiguity_nobiran','ambiguity_variants','ambiguity_classes','ambiguity_pages','ambiguity_traversing'}
assert required<=set(n['report_key'] for n in notes)
unit_order={u['id']:i for i,u in enumerate(units)}
notes.sort(key=lambda n:(unit_order[n['unit_id']],n['part'] if n['part'] is not None else 999))
by_unit=collections.defaultdict(list)
for i,n in enumerate(notes,1):n['number']=i;by_unit[n['unit_id']].append(n)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9)
sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.85)

def set_font(style,name,size):
    style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(rf.attrib):
        if key.endswith('Theme'):del rf.attrib[key]
    for key in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+key),name)

normal=doc.styles['Normal'];set_font(normal,'Cambria',11.5)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for name,size in [('Title',23),('Subtitle',12),('Heading 1',18),('Heading 2',14),('Heading 3',12),('Heading 4',11.5)]:
    st=doc.styles[name];set_font(st,'Cambria',size)
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(13 if name.startswith('Heading') else 0)
    st.paragraph_format.space_after=Pt(6)
for name,size in [('Source',8.5),('Table Text',9.5),('Credits',10),('Poem',11.5),('Comment Text',9)]:
    st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH) if name not in doc.styles else doc.styles[name]
    st.base_style=normal;set_font(st,'Cambria',size)
    st.paragraph_format.space_after=Pt(3)
    if name=='Table Text':st.paragraph_format.line_spacing=1.0
    if name=='Source':st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(8)
    if name=='Poem':st.font.italic=True;st.paragraph_format.space_after=Pt(1)
# The stock template contains a blue border under Title; remove all inherited
# paragraph borders rather than carrying the template decoration into a book.
for style in doc.styles:
    for border in list(style.element.xpath('.//w:pBdr')):border.getparent().remove(border)
doc.core_properties.title='ACKS II Введение и глава 1 с редакторскими пометками'
doc.core_properties.author='Перевод ACKS II'
doc.core_properties.subject='Переработанная книга правил Материалы перед главой 1 и Персонажи'
doc.core_properties.comments='Отдельная сборка из базы перевода. Редакторские пояснения в комментариях Word.'
settings=doc.settings.element
view=OxmlElement('w:view');view.set(qn('w:val'),'print');settings.append(view)
rev=OxmlElement('w:revisionView')
for k,v in [('markup','1'),('comments','1'),('insDel','0'),('formatting','0')]:rev.set(qn('w:'+k),v)
settings.append(rev)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run();run.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
value=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';value.append(text);field.append(value);footer._p.append(field)

def plain(text):
    text=re.sub(r'<br\s*/?>','\n',text)
    text=re.sub(r'<[^>]+>','',text)
    return html.unescape(text).replace('\\|','|').replace('\\*','*').replace('\\_','_')

def formatted(p,text,force_bold=False):
    text=plain(text)
    # Explicit emphasis is presentation only; visible wording is untouched.
    for token in re.split(r'(\*\*.*?\*\*|\*[^*]+?\*)',text,flags=re.S):
        if not token:continue
        bold=token.startswith('**') and token.endswith('**')
        italic=not bold and token.startswith('*') and token.endswith('*')
        value=token[2:-2] if bold else token[1:-1] if italic else token
        for piece in re.split(r'(-(?=\d)|(?<=\w)-(?=\w))',value):
            if piece is None or not piece:continue
            run=p.add_run()
            if piece=='-':run._r.append(OxmlElement('w:noBreakHyphen'))
            else:run.text=piece
            run.bold=bold or force_bold;run.italic=italic
    if not p.runs:p.add_run('')

anchors={}
mapped=[]
def attach(uid,paragraphs):
    for n in by_unit.get(uid,[]):
        index=n['part'] if n['part'] is not None else len(paragraphs)-1
        p=paragraphs[index]
        mark=p.add_run('*'+str(n['number']))
        mark.font.superscript=True;mark.font.size=Pt(10);mark.bold=True
        u=unit_map[uid]
        source=f"PDF {u['pdf_page']}"+(f", печатная стр. {u['printed_page']}" if u['printed_page'] is not None else '')
        comment=doc.add_comment(mark,text=f"*{n['number']} {n['category']}\n{n['text']}\n{source}.",author='Codex',initials='РП')
        # Use a compact, readable comment style; do not invent tracked changes.
        for cp in comment.paragraphs:
            cp.style=doc.styles['Comment Text']
        n['comment_id']=comment.comment_id
        anchors[n['number']]=uid

def bookmark(p,uid):
    ident=len(mapped)+1
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(ident));start.set(qn('w:name'),'u_'+uid.replace('.','_'))
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(ident))
    p._p.insert(1 if p._p.pPr is not None else 0,start);p._p.append(end)

def paragraph(u):
    text=u['russian'];uid=u['id'];page=u['pdf_page']
    parts=text.splitlines() if text.startswith('- ') else [text]
    created=[]
    for value in parts:
        if not value.strip():continue
        style='Normal'
        if u['kind']=='heading':
            style='Heading 1' if (page in (5,8,9,14) and u['block_order']==1) else 'Heading 2' if u['heading_level'] in (3,4) else 'Heading 3' if u['heading_level']==5 else 'Heading 4'
            if page<=4:style='Title' if u['block_order']==1 else 'Subtitle'
        elif u['kind']=='verse':style='Poem'
        elif page==4:style='Credits'
        p=doc.add_paragraph(style=style)
        if page<=3 and u['kind']=='heading':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if value.startswith('- '):
            value=value[2:];p.style='List Bullet'
        elif value.startswith('♦ '):
            value=value[2:];p.style='List Bullet'
        elif re.match(r'^\d+\.\s',value):
            p.paragraph_format.left_indent=Cm(.55);p.paragraph_format.first_line_indent=Cm(-.55)
        formatted(p,value)
        if style=='Poem':p.paragraph_format.keep_with_next=True
        created.append(p)
    if uid in ('rr.p0012.b016','rr.p0013.b006'):created[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
    bookmark(created[0],uid)
    attach(uid,created);mapped.append(uid)
    return created

def table(t):
    tid=t['id'];nr=t['row_count'];nc=t['column_count']
    t_units=sorted((u for u in units if u['table_id']==tid),key=lambda u:(u['row_index'],u['column_index']))
    table=doc.add_table(rows=nr,cols=nc)
    table.autofit=False
    width=16.8
    if nc==2:widths=[14.8,2.0] if t['pdf_page'] in (5,6,7) else [4.0,12.8] if tid=='rr.p0010.t01' else [6.0,10.8] if tid=='rr.p0020.t01' else [8.4,8.4]
    elif nc==3:widths=[4.8,6.0,6.0] if tid=='rr.p0013.t02' else [3.2,6.8,6.8]
    elif tid=='rr.p0021.t01':widths=[3.2,2.4]*3
    elif nc==6:widths=[3.3]+[2.7]*5
    elif nc==7:widths=[2.7]+[2.22]*5+[3.0]
    else:widths=[width/nc]*nc
    for column,cw in zip(table.columns,widths):column.width=Cm(cw)
    # Borders and padding are explicit rather than theme-dependent.
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
    props.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for side,value in [('top',55),('bottom',55),('left',85),('right',85)]:
        edge=OxmlElement('w:'+side);edge.set(qn('w:w'),str(value));edge.set(qn('w:type'),'dxa');margin.append(edge)
    props.append(margin)
    for row in table.rows:
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
    # Merge ranges first, then write exactly once into the surviving cell.
    for u in t_units:
        ri,ci=u['row_index'],u['column_index']
        if u['rowspan']>1 or u['colspan']>1:
            table.cell(ri,ci).merge(table.cell(ri+u['rowspan']-1,ci+u['colspan']-1))
    merged_title=bool(t_units and t_units[0]['row_index']==0 and t_units[0]['colspan']==nc)
    header_rows={0,1} if merged_title else {0}
    for ri in header_rows:
        repeat=OxmlElement('w:tblHeader');table.rows[ri]._tr.get_or_add_trPr().append(repeat)
    for u in t_units:
        ri,ci=u['row_index'],u['column_index']
        cell=table.cell(ri,ci);cell.width=Cm(sum(widths[ci:ci+u['colspan']]))
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0];p.style='Table Text'
        p.paragraph_format.keep_with_next=False;p.paragraph_format.space_after=Pt(0)
        header=ri in header_rows
        if header:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
        elif tid in ('rr.p0013.t01','rr.p0013.t02') and ri%2==0:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FAFAFA');cell._tc.get_or_add_tcPr().append(shade)
        value=plain(u['russian'])
        if tid=='rr.p0022.t01':value=value.replace('длину/высоту','длину / высоту')
        if ri not in header_rows and re.fullmatch(r'[\d\s+−–—.,%/d!()\-]+',value):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if t['pdf_page'] in (5,6,7) and ci==1:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        if header and t['pdf_page'] not in (5,6,7):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        formatted(p,value,force_bold=header or '**' in u['english'])
        if tid=='rr.p0022.t01':
            for run in p.runs:run.font.size=Pt(9)
        bookmark(p,u['id']);attach(u['id'],[p]);mapped.append(u['id'])
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=Pt(3);spacer.add_run().font.size=Pt(3)
    return table

doc.add_paragraph('ACKS II Введение и глава 1',style='Title')
doc.add_paragraph('Материалы перед первой главой и Персонажи',style='Subtitle')
doc.add_paragraph('Звёздочки отмечают места редакторских исправлений. Пояснения помещены в комментарии на полях Word. Пометка «Вопрос к словарю» или «Вопрос к оригиналу» означает, что соответствующий текст сохранён и вопрос остаётся открытым.')
doc.add_paragraph('Номера страниц в оглавлении и внутренних ссылках относятся к оригинальной книге. У каждого исходного блока отдельно указана физическая страница PDF. Эта сборка содержит текст до конца первой главы.')

seen_tables=set()
for page in range(1,24):
    page_units=[u for u in units if u['pdf_page']==page]
    if page in (2,3,4,5,8,14):doc.add_page_break()
    printed=page_units[0]['printed_page']
    label=f'Оригинал PDF {page}'+(f'   Печатная страница {printed}' if printed is not None else '')
    p=doc.add_paragraph(label,style='Source')
    for u in page_units:
        if page==4 and u['block_order'] in (1,2,3):
            if u['block_order']==1:
                parts=[x for x in page_units if x['block_order'] in (1,2,3)]
                title=doc.add_paragraph(style='Title')
                for index,part in enumerate(parts):
                    if index:title.add_run(' ')
                    formatted(title,part['russian'])
                    bookmark(title,part['id']);mapped.append(part['id'])
            continue
        if u['table_id']:
            if u['table_id'] not in seen_tables:
                table(tables[u['table_id']]);seen_tables.add(u['table_id'])
        else:paragraph(u)

assert collections.Counter(mapped)==collections.Counter(unit_map.keys()), 'Missing or duplicate translation unit'
assert len(anchors)==len(notes)
doc.save(FILE)
db.close()

# Visible text is preserved unit by unit; stars and comments are additions.
with ZipFile(FILE) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    comments=etree.fromstring(z.read('word/comments.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    starts=xml.xpath('//w:commentRangeStart/@w:id',namespaces=ns)
    ends=xml.xpath('//w:commentRangeEnd/@w:id',namespaces=ns)
    refs=xml.xpath('//w:commentReference/@w:id',namespaces=ns)
    ids=comments.xpath('//w:comment/@w:id',namespaces=ns)
    assert collections.Counter(starts)==collections.Counter(ids)==collections.Counter(ends)==collections.Counter(refs)
    assert len(ids)==len(notes)
    assert b'relationships/comments' in z.read('word/_rels/document.xml.rels')
    assert b'/word/comments.xml' in z.read('[Content_Types].xml')
    assert len(xml.xpath('//w:bookmarkStart[starts-with(@w:name,"u_")]',namespaces=ns))==1689
(WORK/'annotations.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2),encoding='utf8')
(WORK/'assembly_manifest.json').write_text(json.dumps({'output':str(FILE),'units':mapped,'tables':sorted(seen_tables),'comments':len(notes),'report_items_covered':sorted(required)},ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'file':str(FILE),'units':len(mapped),'tables':len(seen_tables),'comments':len(notes)},ensure_ascii=False,indent=2))
